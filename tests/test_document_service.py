"""
文件名：test_document_service.py
介绍：文档读写服务层测试。这层原先埋在 streamlit_file.py 里，和 UI 混在一起无法单测；
      抽出来之后，四种格式的读取分支与失败兜底都能直接覆盖。
"""
import os
import subprocess
import sys
from pathlib import Path

import pytest

from tools.document_service import (
    read_document, get_document_list, delete_document, extract_documents,
    is_store_empty, open_in_file_manager,
)
from tools.pdf_writer import find_cjk_font, write_pdf


class TestExtractDocuments:
    """抽取层是「索引」和「界面预览」共用的唯一入口，行为必须钉死。"""

    def test_source_metadata_equals_exact_path(self, tmp_path):
        """增量索引的比对和删除都依赖这个字段，路径格式不能被改写。"""
        f = tmp_path / "a.txt"
        f.write_text("内容", encoding="utf-8")

        docs = extract_documents(str(f))

        assert len(docs) == 1
        assert docs[0].metadata["source"] == str(f)

    def test_reads_utf8_with_bom(self, tmp_path):
        """Windows 记事本另存为 UTF-8 会带 BOM，直接按 utf-8 读会在正文开头留一个 \\ufeff。"""
        f = tmp_path / "bom.txt"
        f.write_bytes("中文内容".encode("utf-8-sig"))

        content = extract_documents(str(f))[0].page_content

        assert "中文内容" in content
        assert not content.startswith("﻿")

    def test_reads_gbk_encoded_chinese_txt(self, tmp_path):
        """国内导出的中文 txt 大量是 GBK/GB18030，按 UTF-8 硬读会直接抛异常。"""
        f = tmp_path / "gbk.txt"
        f.write_bytes("这是国标编码的中文内容".encode("gb18030"))

        content = extract_documents(str(f))[0].page_content

        assert "这是国标编码的中文内容" in content

    def test_crlf_is_normalized(self, tmp_path):
        """Windows 的 CRLF 必须归一化：切分器的 "\\n\\n" 分隔符匹配不上 "\\r\\n\\r\\n"，
        段落边界会被忽略，chunk 质量下降。"""
        f = tmp_path / "crlf.txt"
        f.write_bytes("第一段\r\n\r\n第二段\r\n".encode("utf-8"))

        content = extract_documents(str(f))[0].page_content

        assert "\r" not in content
        assert content == "第一段\n\n第二段\n"

    def test_reads_utf16_with_bom(self, tmp_path):
        f = tmp_path / "u16.txt"
        f.write_bytes("UTF-16 编码的中文".encode("utf-16"))

        assert "UTF-16 编码的中文" in extract_documents(str(f))[0].page_content

    def test_docx_includes_table_content(self, tmp_path):
        """只读 paragraphs 会把整张表格丢掉——对知识库来说是一大块内容缺失。"""
        import docx

        d = docx.Document()
        d.add_paragraph("正文段落")
        table = d.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "姓名"
        table.cell(0, 1).text = "职位"
        table.cell(1, 0).text = "张三"
        table.cell(1, 1).text = "工程师"
        f = tmp_path / "t.docx"
        d.save(str(f))

        content = extract_documents(str(f))[0].page_content

        assert "正文段落" in content
        assert "张三" in content and "工程师" in content

    @pytest.mark.skipif(find_cjk_font() is None, reason="系统中没有可用的中文字体")
    def test_pdf_is_split_per_page_with_page_number(self, tmp_path):
        f = tmp_path / "a.pdf"
        write_pdf(f, "第一页的内容")

        docs = extract_documents(str(f))

        assert len(docs) == 1
        assert docs[0].metadata["page"] == 1
        assert docs[0].metadata["source"] == str(f)
        assert "第一页的内容" in docs[0].page_content

    def test_empty_file_yields_no_documents(self, tmp_path):
        f = tmp_path / "e.txt"
        f.write_text("   \n\n  ", encoding="utf-8")
        assert extract_documents(str(f)) == []

    def test_unsupported_extension_raises(self, tmp_path):
        f = tmp_path / "a.xlsx"
        f.write_bytes(b"PK\x03\x04")
        with pytest.raises(ValueError, match="不支持的文件格式"):
            extract_documents(str(f))


class TestReadDocument:
    def test_reads_txt(self, tmp_path):
        f = tmp_path / "a.txt"
        f.write_text("纯文本内容", encoding="utf-8")
        assert read_document(str(f)) == "纯文本内容"

    def test_reads_markdown_as_text(self, tmp_path):
        f = tmp_path / "a.md"
        f.write_text("# 标题\n\n正文", encoding="utf-8")
        assert read_document(str(f)) == "# 标题\n\n正文"

    def test_reads_docx(self, tmp_path):
        import docx

        doc = docx.Document()
        doc.add_paragraph("第一段")
        doc.add_paragraph("")  # 空段落应被过滤
        doc.add_paragraph("第二段")
        f = tmp_path / "a.docx"
        doc.save(str(f))

        content = read_document(str(f))
        assert "第一段" in content and "第二段" in content
        assert content == "第一段\n\n第二段"

    @pytest.mark.skipif(find_cjk_font() is None, reason="系统中没有可用的中文字体")
    def test_reads_pdf_generated_by_own_writer(self, tmp_path):
        """自产自销：用 tools/pdf_writer.py 生成的 PDF，pypdf 要能把中文抽回来。

        抽不回来说明 ToUnicode CMap 有问题 —— 那界面里的「浏览原文」对 PDF 就是坏的。
        """
        f = tmp_path / "a.pdf"
        write_pdf(f, "量子计算入门\n这是正文内容。")
        content = read_document(str(f))
        assert "量子计算入门" in content

    def test_missing_file_returns_message_instead_of_raising(self, tmp_path):
        """界面上点「浏览原文」时文件已被外部删除，不能把整个页面搞崩。"""
        result = read_document(str(tmp_path / "不存在.txt"))
        assert "读取失败" in result

    def test_corrupt_pdf_returns_message_instead_of_raising(self, tmp_path):
        f = tmp_path / "bad.pdf"
        f.write_bytes("这不是一个合法的 PDF 文件".encode("utf-8"))
        result = read_document(str(f))
        assert "读取失败" in result

    def test_unsupported_extension_reports_clearly(self, tmp_path):
        """未知格式不做「当纯文本硬读」的兜底，要给明确提示而不是乱码。"""
        f = tmp_path / "a.xlsx"
        f.write_bytes(b"PK\x03\x04binary garbage")
        result = read_document(str(f))
        assert "不支持的文件格式" in result

    def test_empty_file_reports_clearly(self, tmp_path):
        f = tmp_path / "empty.txt"
        f.write_text("   \n\n  ", encoding="utf-8")
        assert "为空" in read_document(str(f))


class TestOpenInFileManager:
    """打开文档目录。

    这一步由运行 Streamlit 的机器执行（浏览器不允许网页碰本地文件系统），
    所以本地可用、云端会失败——失败必须是「返回 False + 说明」，不能抛异常，
    否则用户点一下按钮整个页面就崩了。
    """

    def test_missing_directory_reports_failure(self, tmp_path):
        opened, detail = open_in_file_manager(tmp_path / "不存在的目录")

        assert opened is False
        assert "目录不存在" in detail

    def test_a_file_is_not_a_valid_target(self, tmp_path):
        """传进来的必须是目录，不能是文件。"""
        f = tmp_path / "a.txt"
        f.write_text("x", encoding="utf-8")

        assert open_in_file_manager(f)[0] is False

    def test_windows_uses_startfile(self, tmp_path, monkeypatch):
        calls = []
        monkeypatch.setattr(sys, "platform", "win32")
        monkeypatch.setattr(os, "startfile", lambda p: calls.append(p), raising=False)

        opened, _ = open_in_file_manager(tmp_path)

        assert opened is True
        assert calls == [str(tmp_path)]

    def test_macos_uses_open(self, tmp_path, monkeypatch):
        calls = []
        monkeypatch.setattr(sys, "platform", "darwin")
        monkeypatch.setattr(subprocess, "Popen", lambda cmd, **kw: calls.append(cmd))

        open_in_file_manager(tmp_path)

        assert calls[0][0] == "open"

    def test_linux_uses_xdg_open(self, tmp_path, monkeypatch):
        calls = []
        monkeypatch.setattr(sys, "platform", "linux")
        monkeypatch.setattr(subprocess, "Popen", lambda cmd, **kw: calls.append(cmd))

        open_in_file_manager(tmp_path)

        assert calls[0][0] == "xdg-open"

    def test_failure_is_returned_not_raised(self, tmp_path, monkeypatch):
        """云端服务器没有图形界面，xdg-open 会失败——不能让它冒到页面上。"""
        def boom(*args, **kwargs):
            raise OSError("no display")

        monkeypatch.setattr(sys, "platform", "linux")
        monkeypatch.setattr(subprocess, "Popen", boom)

        opened, detail = open_in_file_manager(tmp_path)

        assert opened is False
        assert "OSError" in detail


class TestIsStoreEmpty:
    """用于判断「首次运行」：空库且文档目录有文件时，应用会自动建一次索引。"""

    def test_empty_store(self, fake_chroma):
        assert is_store_empty(chroma=fake_chroma) is True

    def test_non_empty_store(self, fake_chroma):
        fake_chroma.docs.append({"source": "a.txt", "content": "x"})
        assert is_store_empty(chroma=fake_chroma) is False


class TestGetDocumentList:
    def test_lists_files_sorted_with_extension_filter(self, tmp_path, fake_chroma):
        (tmp_path / "b.txt").write_text("x", encoding="utf-8")
        (tmp_path / "a.md").write_text("x", encoding="utf-8")
        (tmp_path / "noext").write_text("x", encoding="utf-8")  # 无扩展名，应被忽略
        (tmp_path / "sub").mkdir()  # 子目录，应被忽略

        names = [d["name"] for d in get_document_list(chroma=fake_chroma, data_dir=tmp_path)]
        assert names == ["a.md", "b.txt"]

    def test_chunk_counts_and_indexed_flag(self, tmp_path, fake_chroma):
        indexed = tmp_path / "indexed.txt"
        not_indexed = tmp_path / "not_indexed.txt"
        indexed.write_text("x", encoding="utf-8")
        not_indexed.write_text("x", encoding="utf-8")

        source = os.path.normpath(str(indexed.resolve()))
        for _ in range(3):
            fake_chroma.docs.append({"source": source, "content": ""})

        by_name = {d["name"]: d for d in get_document_list(chroma=fake_chroma, data_dir=tmp_path)}

        assert by_name["indexed.txt"]["chunks"] == 3
        assert by_name["indexed.txt"]["indexed"] is True
        assert by_name["not_indexed.txt"]["chunks"] == 0
        assert by_name["not_indexed.txt"]["indexed"] is False

    def test_empty_metadata_entries_do_not_crash(self, tmp_path, fake_chroma):
        """向量库里可能有历史遗留的、没有 source 的条目。"""
        (tmp_path / "a.txt").write_text("x", encoding="utf-8")
        fake_chroma.docs.append({"source": None, "content": ""})

        docs = get_document_list(chroma=fake_chroma, data_dir=tmp_path)
        assert len(docs) == 1 and docs[0]["chunks"] == 0

    def test_empty_directory_returns_empty_list(self, tmp_path, fake_chroma):
        assert get_document_list(chroma=fake_chroma, data_dir=tmp_path) == []


class TestDeleteDocument:
    def test_removes_chunks_and_file(self, tmp_path, fake_chroma):
        f = tmp_path / "a.txt"
        f.write_text("内容", encoding="utf-8")
        source = os.path.normpath(str(f.resolve()))
        fake_chroma.docs.append({"source": source, "content": "内容"})

        delete_document(str(f), chroma=fake_chroma)

        # 传入的路径必须与向量库 metadata 的 source 一致，否则删不干净
        assert {"source": os.path.normpath(str(f))} in fake_chroma.delete_calls
        assert fake_chroma.docs_of(source) == []
        assert not f.exists()

    def test_missing_file_does_not_raise(self, tmp_path, fake_chroma):
        """文件已被外部删除时，仍应正常清理向量库且不报错。"""
        missing = tmp_path / "gone.txt"
        delete_document(str(missing), chroma=fake_chroma)
        assert not missing.exists()
